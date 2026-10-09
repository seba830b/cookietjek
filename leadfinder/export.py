"""Step 5: write leads.csv and leads.xlsx and print a summary."""

import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ["score", "name", "category", "town", "address", "phone", "website",
           "problems", "pitch_angle", "contacted", "notes", "info"]
WIDTHS = {"score": 7, "name": 32, "category": 13, "town": 14, "address": 30, "phone": 14,
          "website": 34, "problems": 50, "pitch_angle": 70, "contacted": 11, "notes": 30, "info": 60}


def write_csv(leads, path, delimiter=";"):
    # utf-8-sig so Excel shows æ/ø/å correctly when double-clicking the file.
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, delimiter=delimiter, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(leads)


def write_xlsx(leads, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for lead in leads:
        ws.append([lead.get(c, "") for c in COLUMNS])
        row = ws.max_row
        site = lead.get("website")
        if site:
            ws.cell(row, COLUMNS.index("website") + 1).hyperlink = site
        # Colour the score: red = hot lead, yellow = maybe, green = fine site.
        colour = "F8CBAD" if lead["score"] >= 70 else "FFE699" if lead["score"] >= 40 else "C6EFCE"
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=colour)
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[col]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def print_summary(leads):
    no_site = [l for l in leads if not l["website"]]
    social_only = [l for l in no_site if "kun Facebook" in l["problems"]]
    broken = [l for l in leads if l["website"] and l["score"] >= 85]
    print()
    print("=" * 70)
    print(f"Leads i alt:            {len(leads)}")
    print(f"Uden hjemmeside:        {len(no_site)}  (heraf kun sociale medier: {len(social_only)})")
    print(f"Hjemmeside virker ikke: {len(broken)}")
    print(f"Score 70+:              {sum(l['score'] >= 70 for l in leads)}")
    print("-" * 70)
    print("Top 10:")
    for l in leads[:10]:
        print(f"  {l['score']:>3}  {l['name'][:30]:<30} {l['category']:<12} {l['town'][:14]}")
        print(f"       {l['pitch_angle']}")
    print("=" * 70)
